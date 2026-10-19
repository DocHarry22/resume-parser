"""Regression coverage for guest privacy, durable editing and real document exports."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import json
import pytest
import fitz
from docx import Document
from fastapi import FastAPI, UploadFile, HTTPException
from fastapi.testclient import TestClient

from app.core.config import settings, Settings
from app.core.workspaces import WorkspaceMiddleware, RequestBodyLimitMiddleware, sign_workspace, verify_workspace
from app.models.builder_models import ResumeBuilder, ResumeUpdate, ExperienceEntry
from app.models.resume_models import Resume, SkillItem
from app.models.cover_letter_models import CoverLetter, CoverLetterExportRequest
from app.services.builder_service import ResumeBuilderService
from app.services.resume_export import export_resume, download_headers
from app.services.cover_letter_export import generate_pdf
from app.utils.document_loader import load_document


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", str(tmp_path))
    monkeypatch.setattr(settings, "cookie_secure", False)
    monkeypatch.setattr(settings, "workspace_secret", "test-only-" + "s" * 40)
    monkeypatch.setattr(settings, "rate_limit_requests", 1000)
    monkeypatch.setattr(settings, "ai_requests_per_minute", 1000)
    from app.api.builder_routes import router as builder_router
    from app.routers.cover_letter_router import router as cover_router
    from app.api.ai_routes import router as ai_router
    app = FastAPI()
    app.add_middleware(RequestBodyLimitMiddleware)
    app.add_middleware(WorkspaceMiddleware)
    app.include_router(builder_router, prefix="/api")
    app.include_router(cover_router, prefix="/api")
    app.include_router(ai_router, prefix="/api")
    from app.api.workspace_routes import router as workspace_router
    app.include_router(workspace_router, prefix="/api")
    return app


def test_signed_workspace_rejects_tampering_and_expiry():
    secret = b"secret"
    workspace = "w" * 43
    cookie = sign_workspace(workspace, secret, now=100)
    assert verify_workspace(cookie, secret, now=101) == workspace
    assert verify_workspace(cookie[:-1] + ("0" if cookie[-1] != "0" else "1"), secret, now=101) is None
    assert verify_workspace(cookie, secret, now=99) is None
    assert verify_workspace(cookie, secret, now=100 + 366 * 86400) is None


def test_workspace_isolation_persistence_and_delete(app):
    alice, bob = TestClient(app), TestClient(app)
    created = alice.post("/api/builder/create", data={"title": "Private resume"})
    assert created.status_code == 200
    resume_id = created.json()["resume"]["id"]
    cookie = created.headers["set-cookie"]
    assert "HttpOnly" in cookie and "SameSite=lax" in cookie
    assert bob.get(f"/api/builder/{resume_id}").status_code == 404
    assert bob.patch(f"/api/builder/{resume_id}", json={"title": "stolen"}).status_code == 404
    assert bob.get("/api/builder/list/all").json()["count"] == 0
    update = {"summary": {"summary": "A short draft."},
              "experience": [{"company": "Acme", "position": "Engineer", "start_date": "2020",
                              "description": ["Built useful systems"]}]}
    assert alice.patch(f"/api/builder/{resume_id}", json=update).status_code == 200
    # New service instances read persisted data, not a process-local cache.
    reloaded = alice.get(f"/api/builder/{resume_id}").json()["resume"]
    assert reloaded["experience"][0]["company"] == "Acme"
    assert alice.get(f"/api/builder/{resume_id}/export/docx").status_code == 200
    assert alice.delete(f"/api/builder/{resume_id}").status_code == 200
    assert alice.get(f"/api/builder/{resume_id}").status_code == 404


def test_guest_cannot_access_unowned_legacy_documents(app):
    legacy = Path(settings.data_dir) / "resumes"
    legacy.mkdir()
    (legacy / "legacy-id.json").write_text(ResumeBuilder(id="legacy-id", title="Old private data").model_dump_json())
    client = TestClient(app)
    assert client.get("/api/builder/list/all").json()["count"] == 0
    assert client.get("/api/builder/legacy-id").status_code == 404
    assert (legacy / "legacy-id.json").exists()


def test_origin_restriction_blocks_cookie_mutations(app):
    client = TestClient(app)
    response = client.post("/api/builder/create", headers={"Origin": "https://evil.example"})
    assert response.status_code == 403
    assert client.post("/api/builder/create", headers={"Origin": "https://resume.docheng.co.za"}).status_code == 200


def test_cover_letter_workspace_persistence_and_typed_updates(app):
    alice, bob = TestClient(app), TestClient(app)
    response = alice.post("/api/cover-letters", json={"title": "Private letter"})
    assert response.status_code == 200
    letter_id = response.json()["id"]
    assert bob.get("/api/cover-letters").json() == []
    assert bob.get(f"/api/cover-letters/{letter_id}").status_code == 404
    assert alice.put(f"/api/cover-letters/{letter_id}", json={
        "sender": {"full_name": "A Person", "email": "person@example.com"},
        "content": {"opening_paragraph": "Dear team, <b>literal text</b> & experience."}
    }).status_code == 200
    downloaded = alice.post(f"/api/cover-letters/{letter_id}/export", json={"format": "pdf"})
    assert downloaded.status_code == 200
    assert downloaded.content.startswith(b"%PDF")
    persisted = alice.get(f"/api/cover-letters/{letter_id}").json()
    assert persisted["sender"]["full_name"] == "A Person"
    assert alice.delete(f"/api/cover-letters/{letter_id}").status_code == 200
    assert alice.get(f"/api/cover-letters/{letter_id}").status_code == 404


def test_stateless_downloads_and_safe_filename(app):
    client = TestClient(app)
    response = client.post("/api/builder/export/docx", json={
        "title": 'A title"\r\nInjected: yes', "contact": {"full_name": "Jane", "email": ""},
        "custom_sections": {"Awards": ["Original work"]}
    })
    assert response.status_code == 200
    assert "\r" not in response.headers["content-disposition"]
    assert "word/document.xml" in ZipFile(BytesIO(response.content)).namelist()
    response = client.post("/api/cover-letters/export-data?format=pdf", json={
        "sender": {"full_name": "Jane"}, "content": {"opening_paragraph": "Current unsaved letter"}
    })
    assert response.status_code == 200
    with fitz.open(stream=response.content, filetype="pdf") as document:
        assert "Current unsaved letter" in "".join(page.get_text() for page in document)


def test_path_traversal_cannot_read_or_delete(tmp_path):
    service = ResumeBuilderService(tmp_path / "resumes")
    outside = tmp_path / "outside.json"
    outside.write_text('{"secret": true}')
    assert service.get_resume("../outside") is None
    assert service.delete_resume("../outside") is False
    assert service.save_resume(ResumeBuilder(id="../outside")) is False
    assert outside.read_text() == '{"secret": true}'


def test_nested_patch_stays_typed_and_is_durable(tmp_path):
    service = ResumeBuilderService(tmp_path)
    resume = service.create_resume()
    updated = service.update_resume(resume.id, ResumeUpdate(
        summary={"summary": "Short draft"},
        experience=[{"company": "Example", "position": "Dev", "start_date": "2020"}],
        certifications=[{"name": "An actual certificate", "issuer": ""}],
    ))
    assert isinstance(updated.experience[0], ExperienceEntry)
    assert updated.summary.summary == "Short draft"
    loaded = ResumeBuilderService(tmp_path).get_resume(resume.id)
    assert loaded.experience[0].company == "Example"
    assert "An actual certificate" in service.export_to_text(loaded)


def test_import_preserves_missing_contact_and_large_skill_list(tmp_path):
    service = ResumeBuilderService(tmp_path)
    result = service.create_from_parsed(Resume(summary="Short factual summary",
        skills=[SkillItem(name=f"Skill {i}") for i in range(45)]))
    assert result.contact.email == ""
    assert result.contact.full_name == ""
    assert sum(len(group.skills) for group in result.skills) == 45
    assert "email@example.com" not in result.model_dump_json()


def test_pdf_contains_selectable_text_and_paginates():
    resume = ResumeBuilder(
        title="Long resume", contact={"full_name": "Zoë Müller", "email": "zoe@example.com"},
        projects=[{"name": "Real project", "description": "Project supplied by the user.",
                   "highlights": ["Delivered useful work"]}],
        custom_sections={"Research": [f"Evidence record {index} " + "Original supplied content " * 6 for index in range(130)]}
    )
    content, _ = export_resume(resume, "pdf")
    with fitz.open(stream=content, filetype="pdf") as document:
        text = "".join(page.get_text() for page in document)
        assert len(document) > 1
        assert "Real project" in text
        assert "Evidence record 129" in text
        assert all(page.get_text().strip() for page in document)


def test_docx_is_editable_and_preserves_optional_sections():
    resume = ResumeBuilder(title="Original", achievements=["Actual award"],
                           languages=[{"language": "English", "proficiency": "Fluent"}],
                           volunteer=[{"organisation": "Local group", "description": "Helped at events"}],
                           custom_sections={"Memberships": ["Engineering body"]})
    content, _ = export_resume(resume, "docx", "A4")
    document = Document(BytesIO(content))
    paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Engineering body" in paragraphs and "Actual award" in paragraphs
    assert "English" in paragraphs and "Local group" in paragraphs
    assert round(document.sections[0].page_width.mm) == 210
    document.add_paragraph("Editable after export")
    output = BytesIO()
    document.save(output)
    assert len(output.getvalue()) > 0


@pytest.mark.asyncio
async def test_upload_txt_and_signature_validation():
    upload = UploadFile(filename="cv.txt", file=BytesIO(b"Jane Doe\nEngineer\njane@example.com"))
    assert "Jane Doe" in (await load_document(upload)).full_text
    with pytest.raises(HTTPException) as error:
        await load_document(UploadFile(filename="fake.pdf", file=BytesIO(b"malicious HTML")))
    assert error.value.status_code == 422


@pytest.mark.asyncio
async def test_docx_archive_limits():
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<types/>")
        archive.writestr("word/document.xml", b"x" * (2 * 1024 * 1024))
    with pytest.raises(HTTPException) as error:
        await load_document(UploadFile(filename="bomb.docx", file=BytesIO(stream.getvalue())))
    assert error.value.status_code == 422
    assert "unsafe" in error.value.detail


def test_rate_limit_returns_retry_header(app, monkeypatch):
    monkeypatch.setattr(settings, "rate_limit_requests", 1)
    client = TestClient(app)
    assert client.post("/api/builder/create").status_code == 200
    response = client.post("/api/builder/create")
    assert response.status_code == 429 and response.headers["retry-after"] == "60"


def test_cors_wildcard_and_insecure_cross_site_cookies_rejected():
    with pytest.raises(ValueError):
        Settings(cors_origins=["*"], _env_file=None)
    with pytest.raises(ValueError):
        Settings(cookie_secure=False, cookie_samesite="none", _env_file=None)


@pytest.mark.asyncio
async def test_guided_writing_never_invents_experience():
    from app.services.ai_content_generator import generate_cover_letter_content
    output = await generate_cover_letter_content("body", "Engineer",
        job_description="10 years Python leadership required")
    assert output["provider"] == "guided"
    assert all("5+" not in value and "10 years" not in value for value in output["suggestions"])
    assert "Python" not in " ".join(output["suggestions"])


def test_ai_requires_consent_and_configuration(app, monkeypatch):
    monkeypatch.setattr(settings, "ai_api_key", None)
    monkeypatch.setattr(settings, "ai_model", None)
    client = TestClient(app)
    assert client.post("/api/ai/rewrite", json={"text": "I built software."}).status_code == 422
    assert client.post("/api/ai/rewrite", json={"text": "I built software.", "consent": True}).status_code == 503
    assert client.get("/api/ai/status").json()["configured"] is False


@pytest.mark.asyncio
async def test_ai_rejects_invented_metrics(monkeypatch):
    from app.services.writing_provider import rewrite_text
    import app.services.writing_provider as provider
    monkeypatch.setattr(settings, "ai_api_key", "test-not-a-real-key")
    monkeypatch.setattr(settings, "ai_model", "configured-test-model")
    class FakeResponse:
        def raise_for_status(self): pass
        def json(self):
            return {"choices": [{"message": {"content": "Improved performance by 99%."}}]}
    class FakeClient:
        def __init__(self, **kwargs): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *args): pass
        async def post(self, *args, **kwargs): return FakeResponse()
    monkeypatch.setattr(provider.httpx, "AsyncClient", FakeClient)
    with pytest.raises(HTTPException) as error:
        await rewrite_text("Improved performance.", "improve")
    assert error.value.status_code == 422


def test_workspace_export_and_delete_cover_all_guest_documents(app):
    client = TestClient(app)
    client.post("/api/builder/create", data={"title": "My own resume"})
    client.post("/api/cover-letters", json={"title": "My own letter"})
    exported = client.get("/api/workspace/export")
    assert exported.status_code == 200
    assert len(exported.json()["resumes"]) == 1
    assert len(exported.json()["cover_letters"]) == 1
    assert client.delete("/api/workspace").status_code == 200
    assert client.get("/api/workspace").json()["resume_count"] == 0
    assert client.get("/api/cover-letters").json() == []


def test_exports_reject_unsupported_control_characters():
    with pytest.raises(ValueError, match="control characters"):
        export_resume(ResumeBuilder(custom_sections={"Notes": ["Text" + chr(0)]}), "docx")


@pytest.mark.asyncio
async def test_ai_success_returns_provider_text_without_updating_document(monkeypatch):
    from app.services.writing_provider import rewrite_text
    import app.services.writing_provider as provider
    monkeypatch.setattr(settings, "ai_api_key", "test-not-a-real-key")
    monkeypatch.setattr(settings, "ai_model", "configured-test-model")
    observed = {}
    class FakeResponse:
        def raise_for_status(self): pass
        def json(self):
            return {"choices": [{"message": {"content": "Built useful software."}}]}
    class FakeClient:
        def __init__(self, **kwargs):
            observed.update(kwargs)
        async def __aenter__(self): return self
        async def __aexit__(self, *args): pass
        async def post(self, url, **kwargs):
            observed.update(kwargs)
            return FakeResponse()
    monkeypatch.setattr(provider.httpx, "AsyncClient", FakeClient)
    assert await rewrite_text("I built useful software.", "shorten") == "Built useful software."
    assert observed["follow_redirects"] is False
    assert observed["json"]["store"] is False
    assert observed["json"]["max_completion_tokens"] == 1200


def test_import_validation_status_is_preserved(app):
    response = TestClient(app).post("/api/builder/import", files={
        "file": ("fake.pdf", b"HTML instead of PDF", "application/pdf")
    })
    assert response.status_code == 422
    assert "valid PDF" in response.json()["detail"]


def test_chunked_body_limit_cannot_be_bypassed_by_omitting_content_length(app):
    client = TestClient(app)
    def oversized():
        yield b'{"title":"'
        yield b"x" * (1024 * 1024)
        yield b'"}'
    response = client.post("/api/builder/export/json", content=oversized(),
                           headers={"Content-Type": "application/json"})
    assert response.status_code == 413


def test_service_rejects_oversized_saved_document(tmp_path):
    service = ResumeBuilderService(tmp_path)
    with pytest.raises(HTTPException) as error:
        service.create_resume(initial_data={"custom_sections": {"Notes": "x" * (1024 * 1024)}})
    assert error.value.status_code == 413
    assert list(tmp_path.glob("*.json")) == []


def test_atomic_save_retries_a_transient_windows_file_lock(tmp_path, monkeypatch):
    import app.services.builder_service as module
    service = ResumeBuilderService(tmp_path)
    real_replace = module.os.replace
    attempts = []
    def temporarily_locked(source, destination):
        attempts.append(1)
        if len(attempts) == 1:
            raise PermissionError("Simulated transient indexing lock")
        return real_replace(source, destination)
    monkeypatch.setattr(module.os, "replace", temporarily_locked)
    result = service.create_resume(title="Durable despite a brief file lock")
    assert len(attempts) == 2
    assert ResumeBuilderService(tmp_path).get_resume(result.id).title == result.title


def test_failed_atomic_update_keeps_previous_disk_and_cached_content(tmp_path, monkeypatch):
    import app.services.builder_service as module
    service = ResumeBuilderService(tmp_path)
    result = service.create_resume(title="Previously saved")
    def locked(*args):
        raise PermissionError("Simulated persistent disk lock")
    monkeypatch.setattr(module.os, "replace", locked)
    with pytest.raises(OSError):
        service.update_resume(result.id, ResumeUpdate(title="Unsaved mutation"))
    assert service.get_resume(result.id).title == "Previously saved"
    assert ResumeBuilderService(tmp_path).get_resume(result.id).title == "Previously saved"
    assert list(tmp_path.glob("*.tmp")) == []
