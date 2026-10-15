"""Text-based, paginated PDF and editable DOCX resume exports."""
from io import BytesIO
from pathlib import Path
import re
from urllib.parse import quote
from xml.sax.saxutils import escape

from app.models.builder_models import ResumeBuilder


def _flatten(value):
    if isinstance(value, dict):
        for key, item in value.items():
            if item not in (None, "", [], {}):
                yield f"{str(key).replace('_', ' ').title()}: " + "; ".join(_flatten(item))
    elif isinstance(value, list):
        for item in value:
            yield from _flatten(item)
    elif value is not None:
        yield str(value)


def resume_blocks(resume: ResumeBuilder):
    """One canonical reading order across TXT, PDF and DOCX, including custom content."""
    blocks = []
    def text(value, kind="body"):
        if value:
            blocks.append((kind, str(value)))
    def heading(value):
        text(value, "heading")
    if resume.contact:
        contact = resume.contact
        text(contact.full_name.upper(), "name")
        text(" | ".join(v for v in (contact.email, contact.phone, contact.location) if v))
        for label in ("linkedin", "github", "website"):
            value = getattr(contact, label)
            text(f"{label.title()}: {value}" if value else "")
    if resume.summary and resume.summary.summary:
        heading("PROFESSIONAL SUMMARY")
        text(resume.summary.summary)
    if resume.experience:
        heading("EXPERIENCE")
        for item in resume.experience:
            text(" at ".join(v for v in (item.position, item.company) if v), "entry")
            text(item.location)
            text(" - ".join(v for v in (item.start_date, "Present" if item.current else item.end_date) if v))
            for bullet in item.description + item.achievements:
                text(bullet, "bullet")
    if resume.education:
        heading("EDUCATION")
        for item in resume.education:
            text(", ".join(v for v in (item.degree, item.field_of_study, item.institution) if v), "entry")
            text(item.location)
            text(" - ".join(v for v in (item.start_date, item.end_date) if v))
            if item.gpa is not None:
                text(f"GPA: {item.gpa}")
            for honor in item.honors:
                text(honor, "bullet")
    if resume.skills:
        heading("SKILLS")
        for category in resume.skills:
            text(f"{category.category}: {', '.join(category.skills)}")
    if resume.certifications:
        heading("CERTIFICATIONS")
        for item in resume.certifications:
            text(" - ".join(v for v in (item.name, item.issuer, item.issue_date) if v), "entry")
            if item.expiry_date:
                text(f"Expires: {item.expiry_date}")
            if item.credential_id:
                text(f"Credential: {item.credential_id}")
            text(item.credential_url)
    if resume.projects:
        heading("PROJECTS")
        for item in resume.projects:
            text(item.name, "entry")
            text(item.description)
            text(" - ".join(v for v in (item.start_date, item.end_date) if v))
            text(", ".join(item.technologies))
            text(item.url)
            text(item.github)
            for bullet in item.highlights:
                text(bullet, "bullet")
    for name in ("achievements", "languages", "volunteer", "publications"):
        values = getattr(resume, name)
        if values:
            heading(name.upper())
            for value in values:
                for line in _flatten(value):
                    text(line, "bullet")
    for name, value in resume.custom_sections.items():
        heading(str(name).replace("_", " ").upper())
        for line in _flatten(value):
            text(line)
    if not blocks:
        text(resume.title, "name")
    return blocks


def download_headers(title: str, extension: str):
    # Header values must never contain user-controlled quotes, CR/LF or raw Unicode.
    name = re.sub(r"[\x00-\x1f\x7f/\\]", "", title).strip()[:100] or "resume"
    ascii_name = re.sub(r"[^A-Za-z0-9_. -]", "_", name) or "resume"
    return {"Content-Disposition": f'attachment; filename="{ascii_name}.{extension}"; '
            f"filename*=UTF-8''{quote(name + '.' + extension)}",
            "Cache-Control": "no-store"}


def _pdf(resume: ResumeBuilder, page_size: str) -> bytes:
    from reportlab.lib.pagesizes import A4, letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    stream = BytesIO()
    font = "Helvetica"
    # Use a Unicode TrueType font when available on the deployment host.
    for path in (Path(__file__).parent.parent / "assets" / "DejaVuSans.ttf",
                 Path("C:/Windows/Fonts/arial.ttf"),
                 Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")):
        if path.exists():
            if "ResumeUnicode" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ResumeUnicode", str(path)))
            font = "ResumeUnicode"
            break
    values = [value for _, value in resume_blocks(resume)]
    if font == "ResumeUnicode":
        glyphs = pdfmetrics.getFont(font).face.charToGlyph
        unsupported = any(ord(char) > 31 and ord(char) not in glyphs
                          for value in values for char in value)
    else:
        unsupported = any(ord(char) > 255 for value in values for char in value)
    if unsupported:
        raise ValueError("The server PDF font cannot render some characters. Download DOCX or install a suitable Unicode font.")
    styles = getSampleStyleSheet()
    body = ParagraphStyle("ResumeBody", parent=styles["Normal"], fontName=font,
                          fontSize=10, leading=14, spaceAfter=6, splitLongWords=True)
    heading = ParagraphStyle("ResumeSection", parent=body, fontSize=11, leading=16,
                             spaceBefore=12, keepWithNext=True, textColor="#153c5b")
    name = ParagraphStyle("ResumeName", parent=body, fontSize=20, leading=25,
                          spaceAfter=12, keepWithNext=True)
    entry = ParagraphStyle("ResumeEntry", parent=body, fontSize=11, keepWithNext=True)
    story = []
    for kind, value in resume_blocks(resume):
        safe = escape(value).replace("\n", "<br/>")
        if kind == "bullet":
            safe = "• " + safe
        story.append(Paragraph(safe, {"heading": heading, "name": name, "entry": entry}.get(kind, body)))
    document = SimpleDocTemplate(stream, pagesize=letter if page_size == "Letter" else A4,
                                 leftMargin=42, rightMargin=42, topMargin=42, bottomMargin=42,
                                 title=resume.title, author="")
    document.build(story)
    return stream.getvalue()


def _docx(resume: ResumeBuilder, page_size: str) -> bytes:
    from docx import Document
    from docx.shared import Inches, Mm, Pt
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = (
        (Inches(8.5), Inches(11)) if page_size == "Letter" else (Mm(210), Mm(297))
    )
    section.top_margin = section.bottom_margin = Mm(15)
    section.left_margin = section.right_margin = Mm(15)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)
    for kind, value in resume_blocks(resume):
        if kind in {"name", "heading"}:
            document.add_heading(value, level=0 if kind == "name" else 1)
        else:
            paragraph = document.add_paragraph(value, style="List Bullet" if kind == "bullet" else None)
            if kind == "entry":
                paragraph.paragraph_format.keep_with_next = True
                for run in paragraph.runs:
                    run.bold = True
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def export_resume(resume: ResumeBuilder, format: str, page_size: str = "A4"):
    if any(ord(char) < 32 and char not in "\n\r\t"
           for _, value in resume_blocks(resume) for char in value):
        raise ValueError("Remove unsupported control characters before exporting")
    if page_size not in {"A4", "Letter"}:
        raise ValueError("Page size must be A4 or Letter")
    if format == "pdf":
        return _pdf(resume, page_size), "application/pdf"
    if format == "docx":
        return _docx(resume, page_size), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if format in {"txt", "text"}:
        text = "\n".join(("• " if kind == "bullet" else "") + value
                         for kind, value in resume_blocks(resume))
        return text.encode("utf-8"), "text/plain; charset=utf-8"
    if format == "json":
        return resume.model_dump_json(indent=2).encode("utf-8"), "application/json"
    raise ValueError("Supported formats: pdf, docx, txt, json")
